from pptx import Presentation
from pptx.util import Inches
import os


def generate_presentation(slide_images, output_file="outputs/Promotion_Analysis_Report.pptx"):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Promotion Analysis Report"

    overview_slide = prs.slides.add_slide(prs.slide_layouts[1])
    overview_slide.shapes.title.text = 'Overview'
    overview_slide.placeholders[
        1].text = "Impact of promotions on item and store-level sales\nUsing clustering and statistical testing."

    for title, image_path in slide_images:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_picture(image_path, Inches(1), Inches(2), width=Inches(8))

    prs.save(output_file)
    print(f"Presentation saved to {output_file}")
